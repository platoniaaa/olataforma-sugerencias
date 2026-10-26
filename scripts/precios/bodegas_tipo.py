"""Carga la clasificación de bodegas (REAL / VIRT / ELIM / NODISP) a la plataforma.

El ERP mezcla bodegas físicas con bodegas de proceso —dañados, devolución,
scrap, PE por regularizar, importación—. Desde el 20-10-2026 la lista de precios
cuenta el stock de TODAS las bodegas MENOS las **NODISP** (no disponibles):
dañados (con sus tres grafías), devolución, scrap y la bodega sin nombre. REAL,
VIRT y ELIM solo describen la bodega; lo único que saca stock es NODISP.

El Excel lo mantiene Abastecimiento:

    Bodega              | TIPO BOD
    BODEGA SCRAP        | NODISP
    LINDEROS            | REAL
    PE X REGULARIZAR    | VIRT
    CASA MATRIZ         | ELIM

"NO DISPONIBLE" (y variantes como "no_disp") se acepta como alias de NODISP.

LA TABLA DE PRODUCCIÓN ES EL MAESTRO, no un Excel suelto: el archivo de origen
(`1 Tipo bodega.xlsx`, descarga del 08-09) no tiene NODISP y recargarlo volvería a
contar dañados y devolución. Para editar, baja primero la tabla de producción:

    python bodegas_tipo.py --exportar bodegas_actuales.xlsx --env "ruta\\al\\.env"

Eso deja un Excel con las bodegas clasificadas hoy y, al final, las que tienen
stock y nadie clasificó (sin tipo). Se edita y se vuelve a subir.

Uso:

    python bodegas_tipo.py "C:\\ruta\\bodegas.xlsx" --env "ruta\\al\\.env"
    python bodegas_tipo.py archivo.xlsx --api http://localhost:8000 --email x --password y
    python bodegas_tipo.py archivo.xlsx --env ... --simular   # no escribe nada
    python bodegas_tipo.py archivo.xlsx --env ... --sin-no-disponibles
                                  # permite una carga sin ninguna NODISP (rara vez)

Antes de subir compara contra el stock que la plataforma tiene hoy y avisa de
las bodegas que aparecen con stock y no están en el Excel: esas siguen contando
(salvo que su nombre diga dañados, scrap o devolución), pero conviene saber que
existen. El cambio NO recalcula nada: se aplica en el próximo recálculo que relea
el stock (el del motor o el botón Recalcular de /precios).
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import unicodedata
import urllib.request

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from push_precios import CTX  # noqa: E402  (el arreglo de TLS de la red de Curifor)

TIPOS = {"REAL", "VIRT", "ELIM", "NODISP"}
# Como la gente escribe "no disponible"; se compara sobre la clave (sin espacios,
# guiones ni tildes). Igual que `precios_service._ALIAS_TIPO_BODEGA`.
ALIAS = {"NODISPONIBLE": "NODISP", "NODISPONIBLES": "NODISP", "NODISP": "NODISP"}


def clave(nombre) -> str:
    """Igual que `precios_service.clave_bodega`: el nombre sin espacios ni signos.

    "CHILLAN 2" y "CHILLAN2" son la misma bodega; comparando el texto crudo
    quedaban 1,4 millones de unidades sin clasificar.
    """
    txt = str(nombre if nombre is not None else "").upper()
    txt = "".join(c for c in unicodedata.normalize("NFKD", txt) if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]", "", txt)


def normalizar_tipo(valor) -> str:
    """El tipo como lo escribió la persona, normalizado ('' si vino vacío). Un valor
    que no es ninguno de los cuatro (ni alias de NODISP) vuelve tal cual."""
    txt = str(valor if valor is not None else "").strip().upper()
    if txt in TIPOS:
        return txt
    return ALIAS.get(clave(txt), txt)


def leer_excel(ruta: str) -> tuple[list[dict], list[str], list[str]]:
    """(filas, sin_tipo, invalidas).

    - `sin_tipo`: bodegas con nombre y el tipo en blanco. Se omiten (quedan sin
      clasificar) y se avisan.
    - `invalidas`: filas con un tipo que no existe ('NO DISPONBLE', 'FISICA'...). Antes
      se saltaban en silencio, y una NODISP mal escrita dejaba dañados contando;
      ahora frenan la carga.
    """
    wb = load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas, vistas = [], set()
    sin_tipo: list[str] = []
    invalidas: list[str] = []
    for i, fila in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not fila or len(fila) < 2:
            continue
        bodega, tipo = fila[0], normalizar_tipo(fila[1])
        nombre = str(bodega or "").strip()
        if not tipo:
            # Una fila totalmente en blanco no dice nada; una con nombre y sin tipo,
            # si (la exportacion deja asi las bodegas que nadie clasifico).
            if nombre:
                sin_tipo.append(nombre)
            continue
        if tipo not in TIPOS:
            invalidas.append(f"fila {i + 1}: {nombre or '(sin nombre)'} -> {tipo}")
            continue
        k = clave(bodega)
        if k in vistas:
            continue
        vistas.add(k)
        filas.append({"bodega": nombre, "tipo": tipo})
    return filas, sin_tipo, invalidas


def exportar_excel(ruta: str, actuales: list[dict]) -> int:
    """Baja la tabla de producción a un Excel que `leer_excel` vuelve a leer tal
    cual (la tercera columna es solo informativa). Las bodegas con stock que nadie
    clasificó van al final, con el tipo en blanco."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Bodegas"
    ws.append(["Bodega", "TIPO BOD", "Informativo (no se carga)"])
    for f in sorted(actuales, key=lambda f: (bool(f.get("sin_clasificar")), str(f.get("bodega") or ""))):
        if f.get("sin_clasificar"):
            nota = "SIN CLASIFICAR: hoy " + ("cuenta" if f.get("cuenta") else "no cuenta")
            ws.append([f.get("bodega") or None, None, nota])
        else:
            nota = ("cuenta" if f.get("cuenta") else "no cuenta") + (
                "" if f.get("con_stock") else " (hoy sin stock)")
            ws.append([f.get("bodega") or None, f.get("tipo"), nota])
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 40
    wb.save(ruta)
    return len(actuales)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("excel", nargs="?")
    ap.add_argument("--env")
    ap.add_argument("--api")
    ap.add_argument("--email")
    ap.add_argument("--password")
    ap.add_argument("--simular", action="store_true", help="no escribe nada")
    ap.add_argument("--exportar", metavar="ARCHIVO.xlsx",
                    help="baja la clasificación de producción a un Excel (base para editar) y termina")
    ap.add_argument("--sin-no-disponibles", action="store_true",
                    help="permite cargar una clasificación sin ninguna NODISP")
    a = ap.parse_args()

    if a.exportar and a.excel:
        ap.error("--exportar baja la tabla de producción; no se combina con un Excel para cargar")
    if not a.exportar and not a.excel:
        ap.error("falta el Excel para cargar (o --exportar ARCHIVO.xlsx para bajar el actual)")

    base, email, password = a.api, a.email, a.password
    if a.env:
        with open(a.env, encoding="utf-8") as f:
            for linea in f:
                if "=" in linea and not linea.strip().startswith("#"):
                    k, _, v = linea.partition("=")
                    v = v.strip().strip('"')
                    if k.strip() == "PLATAFORMA_API_URL":
                        base = base or v
                    elif k.strip() == "PLATAFORMA_EMAIL":
                        email = email or v
                    elif k.strip() == "PLATAFORMA_PASSWORD":
                        password = password or v
    if not (base and email and password):
        print("Faltan --api/--email/--password o un --env que los traiga", file=sys.stderr)
        return 1
    base = base.rstrip("/")

    filas: list[dict] = []
    if a.excel:
        filas, sin_tipo, invalidas = leer_excel(a.excel)
        if invalidas:
            print(f"{len(invalidas)} fila(s) con un tipo que no existe (se espera REAL, VIRT, ELIM "
                  "o NODISP / 'NO DISPONIBLE'); no se sube nada:", file=sys.stderr)
            for linea in invalidas[:20]:
                print("  " + linea, file=sys.stderr)
            return 1
        if sin_tipo:
            print(f"AVISO: {len(sin_tipo)} bodega(s) sin tipo en el Excel; se omiten y quedan sin "
                  f"clasificar: {', '.join(sin_tipo[:10])}" + (" ..." if len(sin_tipo) > 10 else ""))
        por_tipo: dict[str, int] = {}
        for f in filas:
            por_tipo[f["tipo"]] = por_tipo.get(f["tipo"], 0) + 1
        print(f"Excel: {len(filas)} bodegas  {por_tipo}")
        if not filas:
            print("El Excel no trajo ninguna bodega con tipo REAL/VIRT/ELIM/NODISP", file=sys.stderr)
            return 1
        if not por_tipo.get("NODISP") and not a.sin_no_disponibles:
            print("El Excel no trae ninguna bodega NODISP: recargarlo volvería a contar dañados y "
                  "devolución. Márcalas como NODISP, o usa --sin-no-disponibles si es a propósito.",
                  file=sys.stderr)
            return 1

    def pedir(ruta, metodo="GET", cuerpo=None, token=None):
        datos = json.dumps(cuerpo).encode() if cuerpo is not None else None
        req = urllib.request.Request(base + ruta, data=datos, method=metodo)
        if token:
            req.add_header("Authorization", "Bearer " + token)
        if datos:
            req.add_header("Content-Type", "application/json")
        with urllib.request.urlopen(req, timeout=300, context=CTX) as r:
            return json.load(r)

    token = pedir("/api/auth/login", "POST", {"email": email, "password": password})["token"]

    if a.exportar:
        n = exportar_excel(a.exportar, pedir("/api/precios/bodegas", token=token))
        print(f"Exportadas {n} bodegas de producción a {a.exportar}")
        print("Edítalo y súbelo con: python bodegas_tipo.py ARCHIVO --env ...")
        return 0

    # Que bodegas ve la plataforma hoy con stock, para avisar de las que faltan.
    conocidas = {clave(f["bodega"]) for f in filas}
    try:
        actuales = pedir("/api/precios/bodegas", token=token)
        con_stock = {clave(b["bodega"]) for b in actuales if b.get("con_stock")}
        faltan = con_stock - conocidas
        if faltan:
            print(f"AVISO: {len(faltan)} bodega(s) con stock no estan en el Excel. "
                  "Siguen contando (salvo que el nombre diga dañados, scrap o devolución), "
                  "pero conviene clasificarlas.")
    except Exception as e:  # noqa: BLE001 - es un aviso, no puede frenar la carga
        print(f"(no se pudo comparar contra el stock: {e})")

    if a.simular:
        print("--simular: no se escribio nada.")
        return 0

    cuerpo = {"filas": filas}
    if a.sin_no_disponibles:
        cuerpo["sin_no_disponibles"] = True
    r = pedir("/api/admin/precios/bodegas-tipo", "POST", cuerpo, token)
    print(f"Cargadas: {r.get('bodegas')} bodegas  {r.get('por_tipo')}")
    if r.get("no_disponibles"):
        print(f"No disponibles (no cuentan): {', '.join(r['no_disponibles'])}")
    if r.get("sin_clasificar_con_stock"):
        print(f"AVISO: con stock y sin clasificar: {', '.join(r['sin_clasificar_con_stock'])}")
    print("El cambio no recalcula nada: se aplica en el próximo recálculo que relea el stock "
          "(el del motor, o el botón Recalcular de /precios).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
