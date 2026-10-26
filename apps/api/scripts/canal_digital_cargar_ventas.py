"""Lee los respaldos anuales de ventas (Excel) y arma la venta semanal del canal.

Es lo que `venta_historica` no trae: tipo de venta y fecha del dia. La regla vive en
`src.canal_digital.ventas_semanales` (la misma que usa la plataforma); aqui solo se leen
los Excel y se entrega el resultado.

Por defecto NO publica nada: escribe un CSV con las filas y muestra el resumen. Para
publicar, el motor llama a POST /api/canal-digital/ventas como admin (variables
CANAL_DIGITAL_API y CANAL_DIGITAL_TOKEN, nunca por argumento) con --publicar.

    python scripts/canal_digital_cargar_ventas.py --carpeta "<...>/Bases de datos/Ventas" \
        --anios 2025 2026 --salida ventas_canal.csv
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import pathlib
import sys
import urllib.request

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))

from src.canal_digital.ventas_semanales import (  # noqa: E402
    TIPOS_VENTA_DEFECTO, LineaVenta, agregar, indice_originales, normalizar_codigo, normalizar_numero,
)

COLS = {
    "periodo": "Periodo", "tipo_docto": "tipoDocto", "numero": "Numero", "fecha": "Fecha",
    "tipo_venta": "Tipo-Venta", "producto": "Producto", "cantidad": "Cantidad", "costo": "Costo Neto",
    "doc_ref": "Documento Referencia", "nro_doc_ref": "Nro Documento Referencia",
}


def _num(v) -> float:
    if v is None or isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(".", "").replace(",", ".").strip())
    except ValueError:
        return 0.0


def _fecha(v) -> dt.date | None:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):  # serial de Excel
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))).date()
    t = str(v or "").strip()
    for f in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(t[:10], f).date()
        except ValueError:
            pass
    return None


def leer_archivo(ruta: pathlib.Path, tipos_venta) -> tuple[list[LineaVenta], dict]:
    """Las lineas del canal de un respaldo. Lee en modo read_only (no carga el libro entero)."""
    import openpyxl

    tipos = set(tipos_venta)
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    out: list[LineaVenta] = []
    resumen = {"filas_leidas": 0, "sin_fecha": 0, "sin_producto": 0}
    try:
        it = wb[wb.sheetnames[0]].iter_rows(values_only=True)
        cab = [str(c).strip() if c is not None else "" for c in next(it)]
        faltan = [n for k, n in COLS.items() if k not in ("doc_ref", "nro_doc_ref") and n not in cab]
        if faltan:
            raise ValueError(f"{ruta.name}: faltan columnas {faltan}")
        ix = {k: cab.index(n) if n in cab else None for k, n in COLS.items()}
        for row in it:
            resumen["filas_leidas"] += 1
            if not row:
                continue
            tv = str(row[ix["tipo_venta"]] or "").strip()
            if tv not in tipos:
                continue
            f = _fecha(row[ix["fecha"]])
            if f is None:
                resumen["sin_fecha"] += 1
                continue
            prod = normalizar_codigo(row[ix["producto"]])
            if prod is None:
                resumen["sin_producto"] += 1
                continue
            g = lambda k: row[ix[k]] if ix[k] is not None else None  # noqa: E731
            out.append(LineaVenta(
                periodo=str(g("periodo")).strip().split(".")[0],
                tipo_docto=str(g("tipo_docto") or "").strip(),
                numero=normalizar_numero(g("numero")),
                fecha=f, tipo_venta=tv, producto=prod,
                cantidad=_num(g("cantidad")), costo_neto=_num(g("costo")),
                doc_ref=str(g("doc_ref") or "").strip(), nro_doc_ref=normalizar_numero(g("nro_doc_ref")),
            ))
    finally:
        wb.close()
    return out, resumen


def publicar(filas: list[dict], bodega: str) -> dict:
    base = os.environ.get("CANAL_DIGITAL_API", "").rstrip("/")
    token = os.environ.get("CANAL_DIGITAL_TOKEN", "")
    if not base or not token:
        raise SystemExit("Faltan las variables CANAL_DIGITAL_API y CANAL_DIGITAL_TOKEN (no se pasan por argumento).")
    req = urllib.request.Request(
        f"{base}/api/canal-digital/ventas",
        data=json.dumps({"bodega": bodega, "filas": filas}).encode("utf-8"),
        headers={"Content-Type": "application/json", "Authorization": f"Bearer {token}"}, method="POST",
    )
    with urllib.request.urlopen(req, timeout=300) as r:
        return json.loads(r.read().decode("utf-8"))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--carpeta", required=True, help="carpeta con los respaldos anuales (Ventas)")
    ap.add_argument("--anios", nargs="+", required=True, help="anios a leer, ej. 2025 2026")
    ap.add_argument("--desde-periodo", help="solo cargar periodos >= YYYYMM (los anteriores sirven solo para ubicar originales)")
    ap.add_argument("--bodega", default="430")
    ap.add_argument("--salida", help="CSV con las filas resultantes")
    ap.add_argument("--publicar", action="store_true", help="enviar a la plataforma (admin)")
    a = ap.parse_args()

    carpeta = pathlib.Path(a.carpeta)
    lineas: list[LineaVenta] = []
    for anio in a.anios:
        archivos = [p for p in carpeta.glob("*.xlsx") if anio in p.name and not p.name.startswith("~$")
                    and "Lubricantes" not in p.name and "frontera" not in p.name]
        if not archivos:
            raise SystemExit(f"No hay respaldo de {anio} en {carpeta}")
        for p in archivos:
            print(f"Leyendo {p.name} ...", flush=True)
            l, r = leer_archivo(p, TIPOS_VENTA_DEFECTO)
            print(f"  {r['filas_leidas']:,} filas, {len(l):,} del canal ({r})")
            lineas += l
    originales = indice_originales(lineas, TIPOS_VENTA_DEFECTO)
    periodos = None
    if a.desde_periodo:
        periodos = sorted({l.periodo for l in lineas if l.periodo >= a.desde_periodo})
    res = agregar(lineas, TIPOS_VENTA_DEFECTO, originales, periodos)
    print(f"Filas: {len(res.filas):,}; unidades netas {res.unidades_netas:,.0f}")
    for k, v in sorted(res.estadisticas.items()):
        print(f"  {k}: {v}")
    filas = [f.__dict__ for f in res.filas]
    if a.salida:
        with open(a.salida, "w", newline="", encoding="utf-8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(filas[0].keys()))
            w.writeheader()
            w.writerows(filas)
        print("CSV:", a.salida)
    if a.publicar:
        print("Publicado:", publicar(filas, a.bodega))
    else:
        print("No se publico nada (falta --publicar).")


if __name__ == "__main__":
    main()
