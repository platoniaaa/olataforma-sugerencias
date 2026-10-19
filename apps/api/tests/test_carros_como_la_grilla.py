"""Los carros de compra (y su Excel) piden lo mismo que la grilla del sugerido.

Hasta el 19-10 Compras tenia su propio SELECT y le faltaba lo que la grilla suma en
Python: el minimo InStock y las sugerencias manuales que no caen sobre una fila
visible. Los tres casos son los de la auditoria del 17-10:

(a) repuesto de pauta con minimo 5 y el modelo pide 3: grilla 5, carro 3;
(b) manual en una sucursal sin fila para ese producto, que si tiene fila en otra:
    las unidades no aparecian en ningun carro;
(c) manual sobre una fila con pedir = No: grilla con la manual, carros vacios.

Lo que se compara es la COMPRA NETA de la grilla (sin el traslado desde el CD, que
no se le compra a nadie) contra las unidades de los carros.
"""
import io

from openpyxl import load_workbook

from src.models import RepuestoInstock, Sugerido

FORD = "Ford Motor Company Chile"


def _grilla(client, **params) -> list[dict]:
    r = client.get("/api/sugerido", params={"limit": 5000, **params})
    assert r.status_code == 200, r.text
    return r.json()["items"]


def _carros(client, **params) -> dict:
    r = client.get("/api/compras/carros", params=params)
    assert r.status_code == 200, r.text
    return r.json()


def _neto(fila: dict) -> float:
    neto = fila.get("sugerido_compra_neto")
    return float(neto if neto is not None else (fila.get("total_sugerido_suc") or 0))


def _compra_de_la_grilla(items: list[dict]) -> float:
    """Lo que la grilla manda a comprar: las filas con sucursal que tienen proveedor,
    mas las fabricadas (manual o InStock), que van a 'Sin proveedor asignado' si no lo
    tienen. Las filas del BI sin proveedor no entran a ningun carro (aparte)."""
    return sum(
        _neto(f) for f in items
        if f.get("sucursal_id")
        and (f.get("proveedor") or f.get("origen") in ("manual", "instock"))
    )


def _cantidad(carros: dict, producto: str, proveedor: str | None = None) -> float:
    return sum(
        linea["cantidad"]
        for carro in carros["carros"]
        if proveedor is None or carro["proveedor"] == proveedor
        for linea in carro["lineas"]
        if linea["producto"] == producto
    )


def _pauta(db_session, producto: str, minimo: int = 5, pide: float = 3) -> None:
    db_session.add(Sugerido(
        tenant_id="curifor", producto=producto, descripcion="Filtro de pauta",
        sucursal_id="LINDEROS", nombre_sucursal="Linderos", clasificacion_abc="A",
        proveedor=FORD, pedir="Si", pedir_flag="Si", costo_unitario=1000,
        total_sugerido_suc=pide, sugerido_compra_neto=pide,
        total_valor_sugerido_clp=pide * 1000, stock_activo_suc=0, stock_en_transito_suc=0,
    ))
    db_session.add(RepuestoInstock(
        tenant_id="curifor", producto=producto, part_number=producto.split(" ")[-1],
        marca="FORD", modelos="Ranger", operacion="Filtro", minimo=minimo, activo=True,
    ))
    db_session.commit()


def test_a_el_minimo_instock_llega_al_carro(client, db_session):
    _pauta(db_session, "17 PAUTA1")
    grilla = [f for f in _grilla(client) if f["producto"] == "17 PAUTA1"]
    # 5 en Linderos (3 del modelo + 2 del minimo) y 5 en cada una de las otras tres
    # sucursales con taller, que no tienen fila en el BI y parten de stock 0.
    assert sorted(_neto(f) for f in grilla) == [5, 5, 5, 5]
    carros = _carros(client)
    assert _cantidad(carros, "17 PAUTA1") == 20
    # Las filas fabricadas heredan el proveedor de la fila de Linderos: van al mismo carro.
    assert _cantidad(carros, "17 PAUTA1", FORD) == 20


def test_b_manual_en_una_sucursal_sin_fila(client):
    r = client.post(
        "/api/sugerencias-manuales",
        json={"producto": "20 BXO5W30AA", "sucursal_id": "CURICO", "unidades": 7},
    )
    assert r.status_code in (200, 201), r.text
    grilla = [f for f in _grilla(client) if f["producto"] == "20 BXO5W30AA"]
    assert sorted((f["sucursal_id"], _neto(f)) for f in grilla) == [
        ("CURICO", 7), ("LINDEROS", 6),
    ]
    assert _cantidad(_carros(client), "20 BXO5W30AA", FORD) == 13


def test_c_manual_sobre_una_fila_que_no_se_pedia(client, db_session):
    fila = db_session.query(Sugerido).filter_by(producto="20 BXO5W30AA").one()
    fila.pedir = fila.pedir_flag = "No"
    fila.total_sugerido_suc = fila.sugerido_compra_neto = fila.total_valor_sugerido_clp = 0
    fila.sugerido_traslado = 0
    db_session.commit()
    assert _carros(client)["carros"] == []

    client.post(
        "/api/sugerencias-manuales",
        json={"producto": "20 BXO5W30AA", "sucursal_id": "LINDEROS", "unidades": 5},
    )
    grilla = [f for f in _grilla(client) if f["producto"] == "20 BXO5W30AA"]
    assert [_neto(f) for f in grilla] == [5]
    assert _cantidad(_carros(client), "20 BXO5W30AA", FORD) == 5


def test_el_neto_no_se_duplica_cuando_todo_lo_cubre_el_traslado(client, db_session):
    """Neto 0 porque el CD traslada todo: una manual de 3 es una compra de 3. Antes
    la grilla leia `neto or total` despues de sumar la manual y mostraba 3 + 4 + 3."""
    fila = db_session.query(Sugerido).filter_by(producto="20 BXO5W30AA").one()
    fila.total_sugerido_suc = 4
    fila.sugerido_traslado = 4
    fila.sugerido_compra_neto = 0
    db_session.commit()
    client.post(
        "/api/sugerencias-manuales",
        json={"producto": "20 BXO5W30AA", "sucursal_id": "LINDEROS", "unidades": 3},
    )
    [fila_grilla] = [f for f in _grilla(client) if f["producto"] == "20 BXO5W30AA"]
    assert fila_grilla["total_sugerido_suc"] == 7
    assert fila_grilla["sugerido_compra_neto"] == 3
    assert _cantidad(_carros(client), "20 BXO5W30AA") == 3


def test_el_total_de_los_carros_cuadra_con_la_grilla(client, db_session, en_catalogo):
    """Los tres casos juntos, mas una manual sobre un producto que no esta en el
    sugerido y una fila del BI sin proveedor (que sigue sin entrar a ningun carro)."""
    _pauta(db_session, "17 PAUTA1")
    db_session.add(Sugerido(
        tenant_id="curifor", producto="25 NOPEDIR", descripcion="Correa",
        sucursal_id="TALCA", nombre_sucursal="Talca", proveedor="REPUESTOS DEL SUR",
        pedir="No", pedir_flag="No", total_sugerido_suc=0, sugerido_compra_neto=0,
        costo_unitario=2000,
    ))
    db_session.add(Sugerido(
        tenant_id="curifor", producto="25 SINPROV", descripcion="Sin proveedor",
        sucursal_id="TALCA", nombre_sucursal="Talca", proveedor=None,
        pedir="Si", pedir_flag="Si", total_sugerido_suc=4, sugerido_compra_neto=4,
    ))
    db_session.commit()
    en_catalogo("99 SOLOMANUAL", costo=500)
    for producto, sucursal, unidades in (
        ("20 BXO5W30AA", "CURICO", 7),
        ("25 NOPEDIR", "TALCA", 5),
        ("99 SOLOMANUAL", "RANCAGUA", 2),
    ):
        r = client.post(
            "/api/sugerencias-manuales",
            json={"producto": producto, "sucursal_id": sucursal, "unidades": unidades},
        )
        assert r.status_code in (200, 201), r.text

    items = _grilla(client)
    carros = _carros(client)
    # 6 + 7 (aceite) + 20 (pauta) + 5 (manual sobre pedir=No) + 2 (manual suelta).
    assert _compra_de_la_grilla(items) == 40
    assert carros["total_unidades"] == 40
    assert _cantidad(carros, "99 SOLOMANUAL", "Sin proveedor asignado") == 2
    assert _cantidad(carros, "25 NOPEDIR", "REPUESTOS DEL SUR") == 5
    assert _cantidad(carros, "25 SINPROV") == 0


def test_filtrar_por_proveedor_no_trae_filas_fabricadas_de_otro(client, db_session, en_catalogo):
    _pauta(db_session, "17 PAUTA1")
    en_catalogo("99 SOLOMANUAL", costo=500)
    client.post(
        "/api/sugerencias-manuales",
        json={"producto": "99 SOLOMANUAL", "sucursal_id": "RANCAGUA", "unidades": 2},
    )
    otro = _carros(client, proveedor="NADIE")
    assert otro["carros"] == []
    ford = _carros(client, proveedor="ford")
    assert [c["proveedor"] for c in ford["carros"]] == [FORD]
    assert _cantidad(ford, "17 PAUTA1") == 20


def test_la_orden_de_compra_en_excel_trae_el_minimo_instock(client, db_session):
    _pauta(db_session, "17 PAUTA1")
    r = client.post(
        "/api/compras/export-excel",
        json={"filtros": {"solo_pedir": True}, "proveedor": FORD},
    )
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    cantidades = {
        ws.cell(row=fila, column=1).value: ws.cell(row=fila, column=4).value
        for fila in range(7, ws.max_row + 1)
    }
    assert cantidades["17 PAUTA1"] == 20
    assert cantidades["20 BXO5W30AA"] == 6
