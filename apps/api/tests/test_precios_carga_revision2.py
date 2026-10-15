"""Carga de precios desde Excel: lo que dejo pendiente la segunda revision (14-10-2026).

La segunda pasada comprobo cada hallazgo de la primera y encontro casos vecinos que las
correcciones no cubrian: conflictos que salian como error, memoria sin tope en xlsx y
CSV, paginas de codigos que corrompen las notas, y mas. Un test por escenario.
"""
import io
import zipfile

import pytest

from src.models import PrecioOverride, PrecioProducto
from src.services import politica_precio_service as pol
from src.services import precios_carga_service as carga
from src.services import precios_service as svc

from .test_precios import _producto, lista_cargada  # noqa: F401  (fixture)
from .test_precios_carga import _a_csv, _bajar, _editar, _ficha, _override, _subir, lista  # noqa: F401
from .test_precios_carga_revision import _csv


# ------------------------------------------------ conflicto antes que error
def test_una_edicion_ajena_que_choca_es_un_conflicto_y_no_bloquea_las_demas_filas(client, lista):
    """Otra persona le pone precio fijo a X despues de bajar el archivo; el archivo le
    edita el factor. La combinacion "precio fijo y factor" la causo ella: antes salia
    como error del archivo y no se cargaba NADA."""
    filas = _bajar(client)
    client.put("/api/precios/71 AAA1/override", json={"precio_fijo": 777})      # otra persona
    filas = _editar(filas, "71 AAA1", Factor="2,5")
    filas = _editar(filas, "13 BBB2", Precio_final="1500")

    v = _subir(client, _a_csv(filas)).json()
    r = _subir(client, _a_csv(filas), aplicar=True)

    assert v["errores_total"] == 0
    assert v["conflictos_total"] == 1 and v["conflictos"][0]["producto"] == "71 AAA1"
    assert v["a_cambiar"] == 1
    assert r.status_code == 200 and r.json()["aplicados"] == 1
    assert _override(lista, "13 BBB2").precio_fijo == 1500
    assert _override(lista, "71 AAA1").precio_fijo == 777


def test_un_error_de_verdad_del_archivo_sigue_bloqueando(client, lista):
    filas = _editar(_bajar(client), "71 AAA1", Factor="0,9")

    assert _subir(client, _a_csv(filas)).json()["errores_total"] == 1


# --------------------------------------------- reintento que se superpone
def test_un_reintento_cuyo_plan_es_anterior_a_la_primera_carga_no_acusa_conflicto(client, lista, monkeypatch):
    """Tras un corte de conexion el servidor termina de procesar la primera peticion
    aunque el cliente se desconecte: el reintento ya habia leido el archivo ANTES de
    que la primera confirmara."""
    filas = _editar(_bajar(client), "71 AAA1", Tipo="Pesado")
    contenido = _a_csv(filas)
    plan_viejo = carga.planificar(lista, "lista.csv", contenido)       # el reintento lee primero...
    assert _subir(client, contenido, aplicar=True).json()["aplicados"] == 1   # ...la primera carga confirma

    monkeypatch.setattr(carga, "planificar", lambda *a, **k: plan_viejo)
    r = _subir(client, contenido, aplicar=True)

    assert r.status_code == 200
    assert r.json()["aplicados"] == 0
    assert r.json()["omitidos"] == {"conflictos": 0, "desconocidos": 0, "sin_referencia": 0}


def test_si_no_se_pueden_leer_los_rubros_no_se_carga_nada(client, lista, monkeypatch):
    """`politica.rubros` atrapa el error y devuelve {}: los productos cuyo tipo sale del
    rubro quedaban sin precio (SIN REVISION) y se confirmaba con un 200."""
    filas = _editar(_bajar(client), "71 AAA1", Obs="nota")
    monkeypatch.setattr(carga.politica, "rubros", lambda db: {})

    r = _subir(client, _a_csv(filas), aplicar=True)

    assert r.status_code == 422 and "política" in r.json()["detail"]
    assert _override(lista, "71 AAA1") is None
    assert _ficha(client, "71 AAA1")["precio_final"] == 17800


# --------------------------------------------------------- referencia
def test_una_referencia_con_el_formato_anterior_se_rechaza_como_danada(client, lista):
    filas = _bajar(client)
    ref = filas[0].index("Referencia")
    for f in filas[1:]:
        f[ref] = "|".join(f[ref].split("|")[:-1])          # una parte menos

    v = _subir(client, _a_csv(filas)).json()

    assert v["errores_total"] == len(filas) - 1 and "dañada" in v["errores"][0]["mensaje"]


def test_la_referencia_no_empieza_nunca_con_un_caracter_de_formula(client, lista):
    client.put("/api/precios/71 AAA1/override", json={"tipo_manual": "=HYPERLINK(\"http://x\")"})
    filas = _bajar(client)
    cab = filas[0]
    por = {f[0]: dict(zip(cab, f)) for f in filas[1:]}

    assert por["71 AAA1"]["Tipo"] == "'=HYPERLINK(\"http://x\")"          # protegido
    assert por["71 AAA1"]["Referencia"][0] in "0123456789abcdef"          # empieza con un hash
    assert _subir(client, _a_csv(filas)).json()["a_cambiar"] == 0         # y vuelve igual


# -------------------------------------------------------- notas y Excel
def _como_excel(texto: str, pagina: str = "cp1252") -> bytes:
    """Lo que Excel escribe al guardar como CSV: un "?" por cada unidad UTF-16 que no
    cabe en la pagina de codigos (un emoji ocupa dos, y queda como "??")."""
    salida = []
    for c in texto:
        try:
            c.encode(pagina)
            salida.append(c)
        except UnicodeEncodeError:
            salida.append("?" * (len(c.encode("utf-16-le")) // 2))
    return "".join(salida).encode(pagina)


@pytest.mark.parametrize("obs", [
    "ok 😀 listo", "A → B ≥ 3", "Привет", "mix é → ñ",
    "🔥 oferta 😀 de verano", "dos 😀🚀 y más",           # emoji: dos "?" cada uno
    "con fino y nbsp",                          # el espacio fino no cabe; el nbsp si
])
def test_una_nota_con_caracteres_fuera_de_cp1252_no_se_sobrescribe_con_signos_de_pregunta(client, lista, obs):
    """Al abrir el CSV en Excel y guardarlo como "CSV" (cp1252), lo que no cabe pasa a "?"
    y la nota se tomaba como editada y se reemplazaba."""
    client.put("/api/precios/71 AAA1/override", json={"obs": obs})
    texto = _a_csv(_bajar(client)).decode("utf-8-sig")

    v = _subir(client, _como_excel(texto)).json()

    assert v["a_cambiar"] == 0 and v["errores_total"] == 0


def test_cambiar_una_nota_que_tenia_emoji_si_se_carga(client, lista):
    """La tolerancia es solo para lo que Excel estropea: una nota cambiada de verdad se carga."""
    client.put("/api/precios/71 AAA1/override", json={"obs": "ok 😀 listo"})
    filas = _editar(_bajar(client), "71 AAA1", Obs="revisada")

    r = _subir(client, _como_excel(_a_csv(filas).decode("utf-8-sig")), aplicar=True)

    assert r.status_code == 200 and r.json()["aplicados"] == 1
    assert _override(lista, "71 AAA1").obs == "revisada"


@pytest.mark.parametrize("pagina", ["cp850", "mac_roman"])
def test_el_csv_de_ms_dos_o_de_mac_se_lee_con_su_pagina_de_codigos(client, lista, pagina):
    """"CSV (MS-DOS)" y "CSV (Macintosh)" guardan la tilde con otro byte: la "i" de
    "Si" quedaba como otro caracter, y las notas se sobrescribian con basura."""
    client.put("/api/precios/71 AAA1/override", json={"obs": "revisión año"})
    texto = _a_csv(_bajar(client)).decode("utf-8-sig")

    v = _subir(client, texto.encode(pagina, errors="replace")).json()

    assert v["errores_total"] == 0 and v["a_cambiar"] == 0


@pytest.mark.parametrize("codificacion", ["utf-16-le", "utf-16-be"])
def test_el_utf16_sin_bom_se_lee(client, lista, codificacion):
    filas = _editar(_bajar(client), "71 AAA1", Tipo="Pesado")
    texto = _csv(filas, sep="\t", codificacion="utf-8").decode("utf-8")

    v = _subir(client, texto.encode(codificacion), nombre="lista.txt").json()

    assert v["a_cambiar"] == 1 and v["errores_total"] == 0


# ------------------------------------------------------------ factor
@pytest.mark.parametrize("margen,celda,cambia", [
    (68, "3,13", False),         # 3,125: Excel redondea la mitad hacia arriba, Python al par
    (12.5, "1,143", False),      # tres decimales
    (12.5, "1,14", False),       # dos decimales
    (12.5, "1,15", True),        # una edicion de verdad
    (9.5, "1,11", False),        # 1,105 -> 1,11
    (35, "1,54", False),
    (35, "1,6", True),
    # El redondeo termina en 0: Excel muestra "1,50" y el valor canonico queda "1.5".
    (33.3, "1,50", False),       # 1,4993
    (33.3, "1,5", False),
    (33.3, "1,499", False),
    (33.3, "1,51", True),
    (9, "1,1", False),           # 1,0989 -> "1,10" -> canonico "1.1"
    (9, "1,10", False),
    (0.4876, "1,00", False),     # 1,0049: "1,00" daba 'Factor: 1 no vale' y frenaba todo el archivo
])
def test_el_factor_redondeado_por_excel_no_es_una_edicion(client, lista, margen, celda, cambia):
    client.put("/api/precios/71 AAA1/override", json={"margen_fijo": margen})

    v = _subir(client, _a_csv(_editar(_bajar(client), "71 AAA1", Factor=celda))).json()

    assert v["a_cambiar"] == (1 if cambia else 0), (margen, celda)


# ----------------------------------------------------------------- tipos
def test_un_tipo_de_la_politica_con_barra_vertical_se_puede_elegir(client, lista):
    for proc in ("Nacional", "Importado"):
        pol.guardar_factores(lista, [{"tipo": "Aire|Acond", "procedencia": proc, "factor": 1.9}], "x")

    r = _subir(client, _a_csv(_editar(_bajar(client), "71 AAA1", Tipo="Aire|Acond")), aplicar=True)

    assert r.status_code == 200, r.text
    assert _override(lista, "71 AAA1").tipo_manual == "Aire|Acond"


# --------------------------------------------------------- tope de memoria
def test_un_csv_con_demasiadas_celdas_se_rechaza(client, lista, monkeypatch):
    monkeypatch.setattr(carga, "MAX_CELDAS", 50)

    r = _subir(client, _a_csv(_bajar(client)))

    assert r.status_code == 422 and "celdas" in r.json()["detail"]


def _xlsx_a_mano(hoja: str, hoja2: str | None = None) -> bytes:
    """Un libro minimo escrito a mano: permite hojas que openpyxl no produce (sin
    dimension, con una referencia dañada, con una celda perdida)."""
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    hojas = [("Hoja1", hoja)] + ([("Hoja2", hoja2)] if hoja2 else [])
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/></Types>')
        z.writestr("xl/workbook.xml",
                   f'<workbook xmlns="{ns}" xmlns:r="{rel}"><sheets>'
                   + "".join(f'<sheet name="{n}" sheetId="{i + 1}" r:id="rId{i + 1}"/>' for i, (n, _) in enumerate(hojas))
                   + "</sheets></workbook>")
        z.writestr("xl/_rels/workbook.xml.rels",
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   + "".join(f'<Relationship Id="rId{i + 1}" Type="{rel}/worksheet" Target="worksheets/sheet{i + 1}.xml"/>'
                             for i in range(len(hojas))) + "</Relationships>")
        for i, (_, xml) in enumerate(hojas):
            z.writestr(f"xl/worksheets/sheet{i + 1}.xml", xml)
    return buf.getvalue()


def _hoja(celdas: str, dimension: str = "") -> str:
    return (f'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">{dimension}'
            f"<sheetData>{celdas}</sheetData></worksheet>")


def _texto_en(ref: str, texto: str = "x") -> str:
    fila = "".join(c for c in ref if c.isdigit())
    return f'<row r="{fila}"><c r="{ref}" t="inlineStr"><is><t>{texto}</t></is></c></row>'


def test_un_xlsx_sin_dimension_con_una_celda_perdida_se_rechaza(client, lista):
    """openpyxl en modo streaming no escribe <dimension>, y el precheck dependia de ella:
    el archivo de 11 KB llegaba a 1 GB de memoria."""
    r = _subir(client, _xlsx_a_mano(_hoja(_texto_en("A1048576"))), nombre="lista.xlsx")

    assert r.status_code == 422 and "primera hoja tiene contenido" in r.json()["detail"]


def test_un_xlsx_con_dimension_falsa_tambien_se_rechaza(client, lista):
    hoja = _hoja(_texto_en("A1048576"), dimension='<dimension ref="A1:I200"/>')

    r = _subir(client, _xlsx_a_mano(hoja), nombre="lista.xlsx")

    assert r.status_code == 422 and "primera hoja tiene contenido" in r.json()["detail"]


def test_un_xlsx_con_una_referencia_de_celda_danada_se_rechaza_sin_abrirlo(client, lista):
    """Con r="A99999999999999999999" calamine pide cientos de GB y ABORTA el proceso."""
    r = _subir(client, _xlsx_a_mano(_hoja(_texto_en("A99999999999999999999"))), nombre="lista.xlsx")

    assert r.status_code == 422 and "referencia dañada" in r.json()["detail"]


def test_una_segunda_hoja_enorme_no_bloquea_la_carga(client, lista):
    """Solo se lee la primera hoja: una de apoyo grande no tiene por que impedirlo."""
    from openpyxl import Workbook

    filas = _editar(_bajar(client), "71 AAA1", Tipo="Pesado")
    wb = Workbook()
    ws = wb.active
    for f in filas:
        ws.append(f)
    wb.create_sheet("apoyo")["A300000"] = "x"
    buf = io.BytesIO()
    wb.save(buf)

    r = _subir(client, buf.getvalue(), nombre="lista.xlsx")

    assert r.status_code == 200 and r.json()["a_cambiar"] == 1


def test_celdas_vacias_con_formato_lejos_de_los_datos_no_bloquean_la_carga(client, lista):
    """Formatear una columna entera deja <c r="E300000" s="1"/> sin valor: calamine no
    la lee, y antes impedia subir el archivo."""
    from openpyxl import Workbook

    filas = _editar(_bajar(client), "71 AAA1", Tipo="Pesado")
    wb = Workbook()
    ws = wb.active
    for f in filas:
        ws.append(f)
    ws["E300000"].number_format = "#,##0"
    buf = io.BytesIO()
    wb.save(buf)

    r = _subir(client, buf.getvalue(), nombre="lista.xlsx")

    assert r.status_code == 200 and r.json()["a_cambiar"] == 1


def test_un_xlsx_con_el_flujo_comprimido_dañado_da_un_mensaje_y_no_un_500(client, lista):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for f in _bajar(client):
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    datos = bytearray(buf.getvalue())
    for i in range(len(datos) // 2, len(datos) // 2 + 40):       # un tramo del medio, con el directorio intacto
        datos[i] ^= 0xFF

    r = _subir(client, bytes(datos), nombre="lista.xlsx")

    assert r.status_code == 422 and "No pude abrir el archivo de Excel" in r.json()["detail"]


def test_un_xlsx_que_descomprime_a_mucho_se_rechaza(client, lista, monkeypatch):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for f in _bajar(client):
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    monkeypatch.setattr(carga, "MAX_DESCOMPRIMIDO", 1000)

    r = _subir(client, buf.getvalue(), nombre="lista.xlsx")

    assert r.status_code == 422 and "demasiado grande" in r.json()["detail"]


def test_el_mensaje_de_la_hoja_demasiado_grande_tiene_su_puntuacion(client, lista):
    r = _subir(client, _xlsx_a_mano(_hoja(_texto_en("A400000"))), nombre="lista.xlsx")

    detalle = r.json()["detail"]
    assert "fila 400.000 y la columna 1, mucho más que la lista." in detalle
    assert "sobrantes, bórralas" in detalle


@pytest.mark.parametrize("bloque", [64, 97, 211, 1000, 4096])
def test_el_escaneo_de_la_hoja_no_se_confunde_con_las_etiquetas_cortadas_por_el_borde_del_bloque(
        client, lista, monkeypatch, bloque):
    """Se lee la hoja por bloques de 1 MB, y una etiqueta <c r="..."> que cae justo en el
    borde llega partida. Con un archivo real de 60 mil filas el precheck lo tomaba por una
    referencia dañada; los archivos chicos de las pruebas nunca cruzaban un borde."""
    from openpyxl import Workbook

    monkeypatch.setattr(carga, "_BLOQUE_XML", bloque)
    filas = _editar(_bajar(client), "71 AAA1", Tipo="Pesado")
    wb = Workbook()
    ws = wb.active
    for f in filas:
        ws.append(f)
    buf = io.BytesIO()
    wb.save(buf)

    r = _subir(client, buf.getvalue(), nombre="lista.xlsx")

    assert r.status_code == 200, r.text
    assert r.json()["a_cambiar"] == 1 and r.json()["errores_total"] == 0


def test_el_escaneo_cuenta_la_ultima_fila_y_columna_aunque_la_etiqueta_caiga_en_un_borde(monkeypatch):
    xml = _hoja("".join(_texto_en(f"{c}{n}") for n in range(1, 40) for c in ("A", "B", "AB")))
    for bloque in (50, 77, 128, 513):
        monkeypatch.setattr(carga, "_BLOQUE_XML", bloque)
        assert carga._escanear_hoja(io.BytesIO(xml.encode("utf-8"))) == (39, 28), bloque
