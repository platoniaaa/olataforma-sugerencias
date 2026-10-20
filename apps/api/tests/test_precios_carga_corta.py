"""Carga por Excel corto (20-10-2026): una plantilla de cinco columnas (Código, Precio fijo,
Desde, Hasta, Obs) para tocar pocos productos sin bajar la lista entera.

A diferencia del archivo completo, el corto NO tiene `Referencia`: no hay una descarga contra la
cual comparar, asi que se aplica TAL CUAL a los codigos que trae y a nadie mas. Lo que estos
tests cuidan:

1. Que el modo se detecte bien (con Referencia = completo; sin ella = corto) y que un archivo
   de la lista completa al que le borraron la Referencia no se cargue como si fuera corto.
2. Que una celda vacia NUNCA quite nada: para quitar hay que escribir QUITAR.
3. Que las reglas del periodo sean las de la ficha, y que los errores bloqueen todo.
4. Que reemplazar un margen o un congelado por un precio fijo se avise y cuente como quita.
5. Que los frenos (quitas, carga masiva, tope de 2.000 filas) y la auditoria sean los de siempre.
6. Que la plantilla se descargue (sin chocar con la ruta `/{producto}`) y se pueda subir vacia.
"""
import io
from datetime import date

import pytest
from fastapi.testclient import TestClient

from src.main import app
from src.models import AuditoriaLog, Notificacion, PrecioCambio, PrecioOverride, PrecioProducto
from src.services import precios_carga_service as carga
from src.services import precios_service as svc
from src.services.auth import requiere_auth

from .test_precios import _producto, lista_cargada  # noqa: F401  (fixture)
from .test_precios_carga import _a_csv, _bajar, _ficha, _override, lista  # noqa: F401
from .test_precios_vigencia import HOY, _dias, _poner, reloj  # noqa: F401  (fixture)

CAB = ["Código", "Precio fijo", "Desde", "Hasta", "Obs"]


# ---------------------------------------------------------------- utilidades
def _dmy(d: date) -> str:
    return d.strftime("%d-%m-%Y")


def _corto(*filas, cab=CAB) -> bytes:
    return _a_csv([list(cab), *[list(f) for f in filas]])


def _post(client, contenido, modo=None, aplicar=False, nombre="corto.csv",
          confirmar_quitas=False, confirmar_masiva=False):
    data = {} if modo is None else {"modo": modo}
    if aplicar:
        data.update(confirmar_quitas=str(confirmar_quitas).lower(),
                    confirmar_masiva=str(confirmar_masiva).lower())
    ruta = "/api/precios/carga" if aplicar else "/api/precios/carga/vista-previa"
    return client.post(ruta, data=data, files={"archivo": (nombre, contenido, "application/octet-stream")})


def _vista(client, contenido, **kw) -> dict:
    r = _post(client, contenido, **kw)
    assert r.status_code == 200, r.text
    return r.json()


def _cargar(client, contenido, **kw):
    return _post(client, contenido, aplicar=True, **kw)


def _errores(v) -> list[str]:
    return [e["mensaje"] for e in v["errores"]]


def _recalcular(db):
    svc.recalcular(db, refrescar_insumos=False)


def _auditorias(db) -> int:
    return db.query(AuditoriaLog).count()


# ============================================================ 1. detectar el modo
def test_con_referencia_es_el_archivo_completo(client, lista, reloj):
    v = _vista(client, _a_csv(_bajar(client)))

    assert v["modo"] == "completo" and v["columnas_ignoradas"] == []
    assert v["a_cambiar"] == 0 and v["errores_total"] == 0


def test_sin_referencia_y_sin_columnas_de_la_lista_es_corto(client, lista, reloj):
    v = _vista(client, _corto(["71 AAA1", "990", "", "", ""]))

    assert v["modo"] == "corto" and v["archivo"] == "corto.csv"
    assert (v["filas"], v["a_cambiar"], v["errores_total"]) == (1, 1, 0)
    # En el corto no hay descarga contra la cual comparar: ni conflictos ni filas sin referencia.
    assert v["conflictos_total"] == v["sin_referencia_total"] == v["archivo_viejo_total"] == 0


def test_un_archivo_sin_referencia_con_columnas_de_la_lista_completa_se_rechaza(client, lista, reloj):
    """Quien borre la Referencia de la lista completa no la sube como corta: cada fila traeria
    el precio de HOY en 'Precio fijo' y se volveria precio fijo."""
    contenido = _a_csv([["Producto", "Costo", "Precio ERP", "Estado", "Precio fijo"],
                        ["71 AAA1", "10000", "15000", "OK", "17800"]])

    for modo in (None, "auto", "corto", "completo"):
        r = _post(client, contenido, modo=modo)
        assert r.status_code == 422, modo
        assert "le falta la Referencia" in r.json()["detail"] and "plantilla corta" in r.json()["detail"]
    r = _cargar(client, contenido)
    assert r.status_code == 422 and _override(lista, "71 AAA1") is None


def test_con_una_o_dos_columnas_de_la_lista_el_archivo_sigue_siendo_corto(client, lista, reloj):
    """El rechazo es con tres o mas: una columna suelta (Costo, Estado) se ignora y se avisa."""
    contenido = _a_csv([["Código", "Precio fijo", "Costo", "Estado"], ["71 AAA1", "990", "10000", "OK"]])

    v = _vista(client, contenido)

    assert v["modo"] == "corto" and v["columnas_ignoradas"] == ["Costo", "Estado"]


def test_el_modo_declarado_tiene_que_coincidir_con_el_archivo(client, lista, reloj):
    completo = _a_csv(_bajar(client))
    corto = _corto(["71 AAA1", "990", "", "", ""])

    a = _post(client, completo, modo="corto")
    b = _post(client, corto, modo="completo")
    c = _post(client, corto, modo="raro")

    assert a.status_code == 422 and "Referencia" in a.json()["detail"]
    assert b.status_code == 422 and "Referencia" in b.json()["detail"]
    assert c.status_code == 422 and "modo" in c.json()["detail"].lower()
    # Declarado o por defecto, cuando coincide pasa.
    assert _vista(client, completo, modo="completo")["modo"] == "completo"
    assert _vista(client, completo, modo="auto")["modo"] == "completo"
    assert _vista(client, corto, modo="corto")["modo"] == "corto"
    assert _vista(client, corto, modo="auto")["modo"] == "corto"
    assert _vista(client, corto, modo="")["modo"] == "corto"            # vacio = auto
    assert _vista(client, corto, modo=" Corto ")["modo"] == "corto"


def test_un_modo_que_contradice_al_archivo_tampoco_se_carga(client, lista, reloj):
    corto = _corto(["71 AAA1", "990", "", "", ""])

    r = _cargar(client, corto, modo="completo")
    s = _cargar(client, _a_csv(_bajar(client)), modo="corto")

    assert r.status_code == 422 and s.status_code == 422
    assert _override(lista, "71 AAA1") is None


def test_un_archivo_sin_nada_reconocible_sigue_diciendo_que_falta_la_referencia(client, lista):
    """El mensaje de siempre (lo prueban otros tests), ahora con la salida del corto."""
    r = _post(client, b"Producto;Tipo\r\n71 A;Liviano\r\n")

    assert r.status_code == 422
    assert "Producto y Referencia" in r.json()["detail"] and "plantilla corta" in r.json()["detail"]


def test_un_corto_sin_precio_fijo_desde_ni_hasta_no_sirve(client, lista):
    contenido = _a_csv([["Código", "Obs"], ["71 AAA1", "solo una nota"]])

    for modo in (None, "corto"):
        r = _post(client, contenido, modo=modo)
        assert r.status_code == 422 and "Precio fijo" in r.json()["detail"], modo


@pytest.mark.parametrize("nombre", ["Código", "Codigo", "CODIGO", "código", "Producto", "SKU", " sku "])
def test_el_codigo_se_encuentra_con_sus_alias(client, lista, reloj, nombre):
    contenido = _corto(["71 AAA1", "990", "", "", ""], cab=[nombre, "Precio fijo", "Desde", "Hasta", "Obs"])

    v = _vista(client, contenido)

    assert v["modo"] == "corto" and v["a_cambiar"] == 1


def test_los_encabezados_pueden_estar_bajo_unas_filas_de_leyenda(client, lista, reloj):
    contenido = _a_csv([["Campaña de octubre", "", "", "", ""], ["", "", "", "", ""], CAB,
                        ["71 AAA1", "990", "", "", ""]])

    v = _vista(client, contenido)

    assert v["modo"] == "corto" and v["filas"] == 1 and v["a_cambiar"] == 1


def test_las_columnas_que_no_son_del_corto_se_ignoran_y_se_listan(client, lista, reloj):
    """Margen fijo % y Congelar no entran al corto en esta etapa: ni se leen ni se aplican."""
    cab = ["Código", "Glosa", "Precio fijo", "Margen fijo %", "Congelar", "Hasta", "Obs", ""]
    fila = ["71 AAA1", "AMORT DEL", "990", "35", "Sí", _dmy(_dias(10)), "", ""]

    v = _vista(client, _corto(fila, cab=cab))
    _cargar(client, _corto(fila, cab=cab))

    assert v["columnas_ignoradas"] == ["Glosa", "Margen fijo %", "Congelar"]    # la vacia no cuenta
    assert v["por_campo"] == {"precio_fijo": 1, "hasta": 1}
    o = _override(lista, "71 AAA1")
    assert (o.precio_fijo, o.margen_fijo, o.congelar) == (990, None, False)


def test_un_csv_guardado_como_ms_dos_se_lee_con_su_pagina_de_codigos(client, lista, reloj):
    """"CSV MS-DOS" guarda la ó de Código en cp850: sin leerla bien no se encuentra el encabezado."""
    contenido = _a_csv([CAB, ["71 AAA1", "990", "", "", "descuento por campaña"]], codificacion="cp850")

    v = _vista(client, contenido)
    _cargar(client, contenido)

    assert v["modo"] == "corto" and v["a_cambiar"] == 1
    assert _override(lista, "71 AAA1").obs == "descuento por campaña"


# ================================================================ 2. fijar un precio
def test_un_precio_fijo_con_periodo_se_ve_antes_y_se_aplica(client, lista, reloj):
    contenido = _corto(["71 AAA1", "990", _dmy(HOY), _dmy(_dias(11)), "campaña"])

    v = _vista(client, contenido)
    assert _override(lista, "71 AAA1") is None              # la vista previa no guarda nada
    r = _cargar(client, contenido)

    (c,) = v["cambios"]
    assert c["producto"] == "71 AAA1" and c["glosa"] == "AMORT DEL"
    assert {x["campo"]: (x["antes"], x["despues"]) for x in c["campos"]} == {
        "Precio fijo": ("(sin precio fijo)", "990"), "Desde": ("(sin Desde)", "20-10-2026"),
        "Hasta": ("(sin vencimiento)", "31-10-2026"), "Obs": ("(vacía)", "campaña")}
    assert (c["precio_antes"], c["precio_despues"]) == (17800, 990)
    assert c["vigencia"] == "Rige del 20-10-2026 al 31-10-2026"
    assert v["efecto"] == {"suben": 0, "bajan": 1, "sin_cambio_de_precio": 0, "con_precio_nuevo": 0}
    assert r.status_code == 200, r.text
    assert r.json()["aplicados"] == 1 and r.json()["modo"] == "corto"
    o = _override(lista, "71 AAA1")
    assert (o.precio_fijo, o.vigente_desde, o.vigente_hasta, o.obs) == (990, HOY, _dias(11), "campaña")
    ficha = _ficha(client, "71 AAA1")
    assert (ficha["precio_final"], ficha["estado"], ficha["vigencia"]) == (990, "FIJO", "vigente")


def test_solo_se_tocan_los_codigos_del_archivo(client, lista, reloj):
    antes = {p: _ficha(client, p)["precio_final"] for p in ("13 BBB2", "71 CCC3")}

    _cargar(client, _corto(["71 AAA1", "990", "", "", ""]))

    assert {p: _ficha(client, p)["precio_final"] for p in ("13 BBB2", "71 CCC3")} == antes
    assert _override(lista, "71 CCC3").congelar is True
    assert _override(lista, "13 BBB2").precio_fijo == 990


def test_el_precio_sube_y_baja_se_cuentan_por_separado(client, lista, reloj):
    v = _vista(client, _corto(["71 AAA1", "990", "", "", ""], ["71 CCC3", "2500", "", "", ""]))

    assert v["efecto"] == {"suben": 1, "bajan": 1, "sin_cambio_de_precio": 0, "con_precio_nuevo": 0}


@pytest.mark.parametrize("texto", ["5.200", "$ 5.200", "5200", "5.200,0", "5200.0", " 5.200 "])
def test_los_precios_se_leen_como_los_escribe_excel_en_chile(client, lista, reloj, texto):
    _cargar(client, _corto(["71 AAA1", texto, "", "", ""]))

    assert _override(lista, "71 AAA1").precio_fijo == 5200


# ==================================================== 3. la celda vacia nunca quita
def test_una_celda_vacia_nunca_quita_nada(client, lista, reloj):
    """13 BBB2 tiene precio fijo 990 y la nota "piso"; el archivo solo cambia la nota."""
    client.put("/api/precios/13 BBB2/override", json={"vigente_hasta": _dias(30).isoformat()})
    contenido = _corto(["13 BBB2", "", "", "", "nueva nota"], ["71 CCC3", "", "", "", ""])

    v = _vista(client, contenido)
    _cargar(client, contenido)

    assert v["a_cambiar"] == 1 and v["sin_cambios"] == 1 and v["quitan_total"] == 0
    o = _override(lista, "13 BBB2")
    assert (o.precio_fijo, o.vigente_hasta, o.obs) == (990, _dias(30), "nueva nota")


def test_la_nota_solo_fija_y_no_se_puede_quitar_con_una_celda_vacia(client, lista, reloj):
    _cargar(client, _corto(["13 BBB2", "", "", "", ""]))

    assert _override(lista, "13 BBB2").obs == "piso"


@pytest.mark.parametrize("celdas,esperado,quitan", [
    (["", "", "QUITAR", ""], (990, -1, None), {"hasta": 1}),
    (["", "", "quitar", ""], (990, -1, None), {"hasta": 1}),
    (["", "", "SIN VENCIMIENTO", ""], (990, -1, None), {"hasta": 1}),
    (["", "Sin vencimiento", "", ""], (990, None, 5), {"desde": 1}),
    (["", "QUITAR", "", ""], (990, None, 5), {"desde": 1}),
    (["QUITAR", "", "", ""], (None, None, None), {"precio_fijo": 1}),     # sin precio no hay periodo
])
def test_quitar_se_escribe_con_la_palabra(client, lista, reloj, celdas, esperado, quitan):
    _poner(lista, "71 AAA1", precio_fijo=990.0, vigente_desde=_dias(-1), vigente_hasta=_dias(5))
    _recalcular(lista)

    v = _vista(client, _corto(["71 AAA1", *celdas]))
    r = _cargar(client, _corto(["71 AAA1", *celdas]))

    assert v["quitan"] == quitan and v["quitan_total"] == 1
    assert r.status_code == 200, r.text
    o = _override(lista, "71 AAA1")
    precio, desde, hasta = esperado
    assert (o.precio_fijo, o.vigente_desde, o.vigente_hasta) == (
        precio, None if desde is None else _dias(desde), None if hasta is None else _dias(hasta))


def test_quitar_el_precio_fijo_devuelve_el_producto_a_la_regla(client, lista, reloj):
    _cargar(client, _corto(["13 BBB2", "QUITAR", "", "", ""]))

    assert _override(lista, "13 BBB2").precio_fijo is None
    assert _ficha(client, "13 BBB2")["estado"] == "SIN STOCK"          # 13 BBB2 no tiene stock


def test_quitar_algo_que_no_existe_no_cambia_nada(client, lista, reloj):
    v = _vista(client, _corto(["71 AAA1", "QUITAR", "QUITAR", "QUITAR", ""]))

    assert v["a_cambiar"] == 0 and v["sin_cambios"] == 1 and v["quitan_total"] == 0


# ========================================================= 4. valores que no valen
@pytest.mark.parametrize("precio,dice", [
    ("0", "mayor que 0"), ("-5", "mayor que 0"), ("0,4", "mayor que 0"),
    ("abc", "no entiendo"), ("1e999", "número"),
])
def test_un_precio_fijo_que_no_vale_es_un_error_que_bloquea_todo(client, lista, reloj, precio, dice):
    contenido = _corto(["71 AAA1", precio, "", "", ""], ["71 CCC3", "2500", "", "", ""])

    v = _vista(client, contenido)
    r = _cargar(client, contenido)

    assert v["errores_total"] == 1 and v["errores"][0]["producto"] == "71 AAA1"
    assert v["errores"][0]["fila"] == 2 and dice in v["errores"][0]["mensaje"]
    assert r.status_code == 422 and "1 error" in r.json()["detail"]
    assert _override(lista, "71 CCC3").precio_fijo is None            # la otra fila tampoco se cargo


@pytest.mark.parametrize("texto,esperado", [
    ("31-10-2026", date(2026, 10, 31)), ("2026-10-31", date(2026, 10, 31)),
    ("31/10/2026", date(2026, 10, 31)), ("31-10-26", date(2026, 10, 31)),
    ("46326", date(2026, 10, 31)), ("2026-10-31 00:00:00", date(2026, 10, 31)),
])
def test_las_fechas_se_leen_como_las_escribe_excel_en_chile(client, lista, reloj, texto, esperado):
    _cargar(client, _corto(["71 AAA1", "990", "", texto, ""]))

    assert _override(lista, "71 AAA1").vigente_hasta == esperado


def test_nunca_se_lee_mes_dia(client, lista, reloj):
    """10/11/2026 es el 10 de noviembre; 10/13/2026 no es una fecha."""
    _cargar(client, _corto(["71 AAA1", "990", "", "10/11/2026", ""]))
    malo = _vista(client, _corto(["71 AAA1", "990", "", "10/13/2026", ""]))

    assert _override(lista, "71 AAA1").vigente_hasta == date(2026, 11, 10)
    assert malo["errores_total"] == 1 and "dd-mm-aaaa" in malo["errores"][0]["mensaje"]


@pytest.mark.parametrize("texto", ["31-oct", "31-02-2026", "mañana", "1e3", "31.10.2026"])
def test_una_fecha_que_no_se_entiende_es_un_error(client, lista, reloj, texto):
    v = _vista(client, _corto(["71 AAA1", "990", "", texto, ""]))

    assert v["errores_total"] == 1 and "Hasta" in v["errores"][0]["mensaje"]
    assert "dd-mm-aaaa" in v["errores"][0]["mensaje"]


def test_las_reglas_del_periodo_son_las_de_la_ficha(client, lista, reloj):
    desde_despues = _vista(client, _corto(["71 AAA1", "990", _dmy(_dias(9)), _dmy(_dias(3)), ""]))
    vencido = _vista(client, _corto(["71 AAA1", "990", "", _dmy(_dias(-1)), ""]))
    lejos = _vista(client, _corto(["71 AAA1", "990", "", "31-10-2206", ""]))

    assert "no puede ser posterior" in _errores(desde_despues)[0]
    assert "ya terminó" in _errores(vencido)[0] and _dmy(_dias(-1)) in _errores(vencido)[0]
    assert "3 años" in _errores(lejos)[0]


def test_hoy_es_el_de_chile_y_el_ultimo_dia_todavia_rige(client, lista, reloj):
    reloj(date(2026, 10, 25))

    hoy = _vista(client, _corto(["71 AAA1", "990", "", "25-10-2026", ""]))
    ayer = _vista(client, _corto(["71 AAA1", "990", "", "24-10-2026", ""]))

    assert hoy["errores_total"] == 0 and hoy["a_cambiar"] == 1
    assert ayer["errores_total"] == 1 and "ya terminó" in ayer["errores"][0]["mensaje"]


def test_solo_el_hasta_sobre_un_producto_sin_decision_es_un_error(client, lista, reloj):
    v = _vista(client, _corto(["71 AAA1", "", "", _dmy(_dias(10)), ""]))

    assert v["errores_total"] == 1
    assert "Precio fijo" in _errores(v)[0]                  # dice que hay que poner tambien el precio
    assert _cargar(client, _corto(["71 AAA1", "", "", _dmy(_dias(10)), ""])).status_code == 422


def test_una_nota_de_mas_de_500_caracteres_es_un_error(client, lista, reloj):
    v = _vista(client, _corto(["71 AAA1", "990", "", "", "x" * 501]))

    assert v["errores_total"] == 1 and "Obs" in _errores(v)[0]


def test_una_fila_con_valores_y_sin_codigo_es_un_error(client, lista, reloj):
    v = _vista(client, _corto(["", "990", "", "", ""], ["71 AAA1", "990", "", "", ""]))

    assert v["errores_total"] == 1 and v["errores"][0]["fila"] == 2
    assert "código" in _errores(v)[0]


def test_las_filas_vacias_no_cuentan(client, lista, reloj):
    v = _vista(client, _corto(["71 AAA1", "990", "", "", ""], ["", "", "", "", ""], ["", "", "", "", ""]))

    assert v["filas"] == 1 and v["errores_total"] == 0


# ============================================================= 5. los codigos
def test_el_codigo_se_recorta_en_los_bordes_y_se_prueba_en_mayusculas(client, lista, reloj):
    v = _vista(client, _corto(["  71 aaa1  ", "990", "", "", ""]))

    assert v["a_cambiar"] == 1 and v["desconocidos_total"] == 0
    assert v["cambios"][0]["producto"] == "71 AAA1"


def test_el_codigo_con_el_espacio_final_del_erp_se_encuentra(client, lista, reloj):
    """Por dentro la lista guarda el codigo sin el espacio; el archivo del ERP lo trae."""
    p = lista.query(PrecioProducto).filter_by(producto="71 AAA1").one()
    p.codigo_erp = "71 AAA1 "
    lista.commit()

    v = _vista(client, _corto(["71 AAA1 ", "990", "", "", ""]))

    assert v["a_cambiar"] == 1 and v["desconocidos_total"] == 0


def test_un_espacio_de_adentro_distinto_es_otro_codigo(client, lista, reloj):
    """"71  AAA1" (dos espacios) no es "71 AAA1": no se colapsan."""
    v = _vista(client, _corto(["71  AAA1", "990", "", "", ""]))

    assert v["desconocidos"] == ["71  AAA1"] and v["a_cambiar"] == 0


def test_dos_codigos_que_solo_difieren_en_un_espacio_de_adentro_son_productos_distintos(client, lista, reloj):
    _producto(lista, "71  DOBLE", procedencia_maestro="NACIONAL")
    _producto(lista, "71 DOBLE", procedencia_maestro="NACIONAL")
    _recalcular(lista)

    r = _cargar(client, _corto(["71  DOBLE", "2500", "", "", ""]))

    assert r.status_code == 200 and r.json()["aplicados"] == 1
    assert _override(lista, "71  DOBLE").precio_fijo == 2500 and _override(lista, "71 DOBLE") is None


def test_un_codigo_repetido_es_un_error(client, lista, reloj):
    v = _vista(client, _corto(["71 AAA1", "990", "", "", ""], ["71 CCC3", "990", "", "", ""],
                              ["71 AAA1", "1500", "", "", ""]))

    assert v["errores_total"] == 1 and "repetido" in _errores(v)[0] and "fila 2" in _errores(v)[0]
    assert v["errores"][0]["fila"] == 4


def test_el_mismo_producto_escrito_en_mayusculas_y_en_minusculas_es_repetido(client, lista, reloj):
    v = _vista(client, _corto(["71 AAA1", "990", "", "", ""], ["71 aaa1", "1500", "", "", ""]))

    assert v["errores_total"] == 1 and "repetido" in _errores(v)[0]


def test_los_desconocidos_se_omiten_y_se_listan(client, lista, reloj):
    contenido = _corto(["99 NOEXISTE", "990", "", "", ""], ["71 AAA1", "990", "", "", ""])

    v = _vista(client, contenido)
    r = _cargar(client, contenido)

    assert v["desconocidos"] == ["99 NOEXISTE"] and v["desconocidos_total"] == 1
    assert v["a_cambiar"] == 1 and v["errores_total"] == 0
    assert r.status_code == 200
    assert r.json()["aplicados"] == 1 and r.json()["omitidos"]["desconocidos"] == 1


def test_un_error_de_valor_bloquea_aunque_haya_filas_buenas_y_desconocidas(client, lista, reloj):
    contenido = _corto(["99 NOEXISTE", "990", "", "", ""], ["71 AAA1", "990", "", "", ""],
                       ["71 CCC3", "abc", "", "", ""])

    r = _cargar(client, contenido)

    assert r.status_code == 422
    assert _override(lista, "71 AAA1") is None


# ========================================================= 6. subir dos veces
def test_subir_dos_veces_el_mismo_archivo_da_sin_cambios(client, lista, reloj):
    contenido = _corto(["71 AAA1", "990", _dmy(HOY), _dmy(_dias(10)), "campaña"])

    primera = _cargar(client, contenido).json()
    v = _vista(client, contenido)
    segunda = _cargar(client, contenido).json()

    assert primera["aplicados"] == 1
    assert (v["a_cambiar"], v["sin_cambios"], v["conflictos_total"]) == (0, 1, 0)
    assert v["sin_cambios_codigos"] == ["71 AAA1"]
    assert segunda["aplicados"] == 0 and segunda["sin_cambios"] == 1
    assert lista.query(AuditoriaLog).filter_by(accion="precio_override").count() == 1


def test_el_corto_se_aplica_contra_lo_que_hay_al_cargar_y_no_da_conflicto(client, lista, reloj):
    """No hay descarga contra la cual comparar: lo que otra persona haya hecho en la ficha
    entre la vista previa y la carga no es un conflicto, el archivo dice lo que dice."""
    contenido = _corto(["71 AAA1", "990", "", "", ""])
    _vista(client, contenido)
    client.put("/api/precios/71 AAA1/override", json={"obs": "otra persona"})

    r = _cargar(client, contenido)

    assert r.status_code == 200 and r.json()["aplicados"] == 1
    assert r.json()["omitidos"] == {"conflictos": 0, "desconocidos": 0, "sin_referencia": 0}
    o = _override(lista, "71 AAA1")
    assert (o.precio_fijo, o.obs) == (990, "otra persona")


# ============================================ 7. reemplazar margen o congelado
def test_un_precio_fijo_sobre_un_margen_lo_reemplaza_y_lo_informa(client, lista, reloj):
    client.put("/api/precios/71 AAA1/override", json={"margen_fijo": 35})
    contenido = _corto(["71 AAA1", "990", "", "", ""])

    v = _vista(client, contenido)
    r = _cargar(client, contenido)

    assert v["reemplazos_total"] == 1
    (re_,) = v["reemplazos"]
    assert re_["producto"] == "71 AAA1" and "Margen fijo 35 %" in re_["decision_anterior"]
    assert re_["editado_por"] == "test@curifor.com" and re_["editado_en"]
    assert v["quitan"] == {"factor": 1} and v["quitan_total"] == 1
    assert r.status_code == 200, r.text
    o = _override(lista, "71 AAA1")
    assert (o.precio_fijo, o.margen_fijo) == (990, None)
    assert _ficha(client, "71 AAA1")["estado"] == "FIJO"


def test_un_precio_fijo_sobre_un_congelado_lo_reemplaza_y_lo_informa(client, lista, reloj):
    v = _vista(client, _corto(["71 CCC3", "2500", "", "", ""]))
    _cargar(client, _corto(["71 CCC3", "2500", "", "", ""]))

    (re_,) = v["reemplazos"]
    assert re_["producto"] == "71 CCC3" and "Congelado" in re_["decision_anterior"]
    assert v["quitan"] == {"congelar": 1}
    o = _override(lista, "71 CCC3")
    assert (o.precio_fijo, o.congelar, o.congelado_precio) == (2500, False, None)


def test_el_reemplazo_dice_si_la_decision_anterior_estaba_vencida(client, lista, reloj):
    _poner(lista, "71 AAA1", margen_fijo=35.0, vigente_desde=_dias(-20), vigente_hasta=_dias(-5))
    _recalcular(lista)

    v = _vista(client, _corto(["71 AAA1", "990", "", _dmy(_dias(10)), ""]))

    assert v["errores_total"] == 0 and v["reemplazos_total"] == 1
    assert "Margen fijo 35 %" in v["reemplazos"][0]["decision_anterior"]
    assert "vencido" in v["reemplazos"][0]["decision_anterior"].lower()


def test_un_precio_fijo_sobre_un_periodo_vencido_pide_un_hasta_nuevo(client, lista, reloj):
    """Hay un solo periodo por producto: el precio nuevo heredaria el Hasta que ya paso, y el
    periodo nacería vencido (lo mismo que rechaza la ficha). Se renueva el Hasta, o se escribe
    QUITAR para dejarlo sin vencimiento."""
    _poner(lista, "71 AAA1", margen_fijo=35.0, vigente_desde=_dias(-20), vigente_hasta=_dias(-5))
    _recalcular(lista)

    sin = _vista(client, _corto(["71 AAA1", "990", "", "", ""]))
    renovado = _vista(client, _corto(["71 AAA1", "990", "", _dmy(_dias(10)), ""]))
    permanente = _vista(client, _corto(["71 AAA1", "990", "", "QUITAR", ""]))

    assert sin["errores_total"] == 1 and "ya terminó" in _errores(sin)[0]
    assert renovado["errores_total"] == 0 and permanente["errores_total"] == 0
    assert permanente["sin_hasta_total"] == 1


def test_un_precio_fijo_sobre_otro_precio_fijo_no_es_un_reemplazo(client, lista, reloj):
    v = _vista(client, _corto(["13 BBB2", "1500", "", "", ""]))

    assert v["reemplazos_total"] == 0 and v["reemplazos"] == [] and v["quitan_total"] == 0


def test_los_reemplazos_cuentan_para_confirmar_quitas(client, lista, reloj, monkeypatch):
    monkeypatch.setattr(carga, "UMBRAL_QUITAS", 1)
    client.put("/api/precios/71 AAA1/override", json={"margen_fijo": 35})
    contenido = _corto(["71 AAA1", "990", "", "", ""])

    sin = _cargar(client, contenido)
    con = _cargar(client, contenido, confirmar_quitas=True)

    assert sin.status_code == 422 and "quita 1 decisiones" in sin.json()["detail"]
    assert con.status_code == 200 and con.json()["aplicados"] == 1


def test_mas_de_200_precios_fijos_pide_confirmar_la_carga_masiva(client, lista, reloj):
    codigos = [f"71 MAS{i:03d}" for i in range(201)]
    for c in codigos:
        _producto(lista, c, procedencia_maestro="NACIONAL")
    contenido = _corto(*[[c, "1500", "", "", ""] for c in codigos])

    v = _vista(client, contenido)
    sin = _cargar(client, contenido)
    nada = lista.query(PrecioOverride).filter(PrecioOverride.producto.like("71 MAS%")).count()
    con = _cargar(client, contenido, confirmar_masiva=True)

    assert (v["fijan_total"], v["umbral_masiva"], v["fijan"]) == (201, 200, {"precio_fijo": 201})
    assert sin.status_code == 422 and "precios fijos" in sin.json()["detail"]
    assert "fila de su código" in sin.json()["detail"]         # el texto propio del corto
    assert nada == 0
    assert con.status_code == 200 and con.json()["aplicados"] == 201


def test_el_umbral_masivo_es_un_mas_de_y_no_un_igual_a(client, lista, reloj, monkeypatch):
    monkeypatch.setattr(carga, "UMBRAL_MASIVA", 2)

    dos = _cargar(client, _corto(["71 AAA1", "990", "", "", ""], ["71 CCC3", "2500", "", "", ""]))
    tres = _cargar(client, _corto(["71 AAA1", "991", "", "", ""], ["71 CCC3", "2501", "", "", ""],
                                  ["13 BBB2", "1500", "", "", ""]))

    assert dos.status_code == 200
    assert tres.status_code == 422


def test_las_fechas_cuentan_para_la_carga_masiva(client, lista, reloj, monkeypatch):
    """Ordenar solo la columna Hasta deja cada fecha en la fila de otro producto."""
    monkeypatch.setattr(carga, "UMBRAL_MASIVA", 1)
    contenido = _corto(["71 AAA1", "990", "", _dmy(_dias(10)), ""])

    v = _vista(client, contenido)

    assert v["fijan"] == {"precio_fijo": 1, "hasta": 1} and v["fijan_total"] == 2
    assert _cargar(client, contenido).status_code == 422


# ============================================================ 8. el periodo
def test_un_desde_futuro_no_cambia_el_precio_de_hoy_y_se_cuenta(client, lista, reloj):
    contenido = _corto(["71 AAA1", "990", _dmy(_dias(5)), _dmy(_dias(30)), ""])

    v = _vista(client, contenido)
    _cargar(client, contenido)

    assert v["programados_total"] == 1
    (c,) = v["cambios"]
    assert c["precio_antes"] == c["precio_despues"] == 17800
    assert c["vigencia"] == "Rige del 25-10-2026 al 19-11-2026"
    ficha = _ficha(client, "71 AAA1")
    assert (ficha["precio_final"], ficha["vigencia"], ficha["precio_fijo"]) == (17800, "programada", 990)


def test_renovar_un_vencido_es_cambiar_solo_el_hasta(client, lista, reloj):
    _poner(lista, "71 AAA1", precio_fijo=990.0, vigente_desde=_dias(-20), vigente_hasta=_dias(-5))
    _recalcular(lista)
    assert _ficha(client, "71 AAA1")["vigencia"] == "vencida"

    v = _vista(client, _corto(["71 AAA1", "", "", _dmy(_dias(30)), ""]))
    _cargar(client, _corto(["71 AAA1", "", "", _dmy(_dias(30)), ""]))

    assert v["errores_total"] == 0 and v["a_cambiar"] == 1
    ficha = _ficha(client, "71 AAA1")
    assert (ficha["vigencia"], ficha["precio_final"], ficha["estado"]) == ("vigente", 990, "FIJO")


def test_un_precio_fijo_sin_hasta_queda_permanente_y_se_avisa(client, lista, reloj):
    contenido = _corto(["71 AAA1", "990", "", "", ""],
                       ["71 CCC3", "2500", "", _dmy(_dias(10)), ""],
                       ["13 BBB2", "", "", "", "solo una nota"])

    v = _vista(client, contenido)

    assert v["sin_hasta_total"] == 1 and v["sin_hasta"] == ["71 AAA1"]
    assert v["errores_total"] == 0 and v["a_cambiar"] == 3


def test_no_se_exige_nota(client, lista, reloj):
    v = _vista(client, _corto(["71 AAA1", "990", "", _dmy(_dias(10)), ""]))

    assert v["errores_total"] == 0 and v["a_cambiar"] == 1


# ======================================================= 9. lo que deja la carga
def test_la_carga_deja_rastro_en_auditoria_y_notificacion(client, lista, reloj):
    _cargar(client, _corto(["71 AAA1", "990", "", _dmy(_dias(10)), "campaña"]))

    por = {a.accion: a for a in lista.query(AuditoriaLog).filter(
        AuditoriaLog.accion.in_(["precio_override", "precios_carga_excel"])).all()}
    ficha = por["precio_override"]
    assert (ficha.producto, ficha.usuario_email, ficha.motivo) == ("71 AAA1", "test@curifor.com", "campaña")
    assert "precio fijo None -> 990" in ficha.detalle and "(carga corta desde Excel)" in ficha.detalle
    assert "(carga desde Excel)" not in ficha.detalle
    resumen = por["precios_carga_excel"]
    assert resumen.usuario_email == "test@curifor.com"
    assert "1 productos" in resumen.detalle and "carga corta" in resumen.detalle.lower()
    notis = [n for n in lista.query(Notificacion).filter_by(tipo="precios").all()
             if "carga corta" in n.titulo.lower()]
    assert len(notis) == 1 and "1 productos" in notis[0].mensaje


def test_la_vista_previa_no_guarda_nada(client, lista, reloj):
    antes = (_auditorias(lista), lista.query(Notificacion).count())

    _vista(client, _corto(["71 AAA1", "990", "", "", "x"], ["71 CCC3", "2500", "", "", ""]))

    assert _override(lista, "71 AAA1") is None and _override(lista, "71 CCC3").precio_fijo is None
    assert (_auditorias(lista), lista.query(Notificacion).count()) == antes
    assert _ficha(client, "71 AAA1")["precio_final"] == 17800


def test_la_carga_no_llena_la_lista_de_cambios_sin_revisar(client, lista, reloj):
    """Son decisiones de una persona, no cambios que alguien tenga que revisar."""
    antes = lista.query(PrecioCambio).count()

    _cargar(client, _corto(["71 AAA1", "990", "", "", ""]))

    assert lista.query(PrecioCambio).count() == antes
    assert _ficha(client, "71 AAA1")["cambios_pendientes"] == 0


def test_la_carga_recalcula_solo_lo_tocado(client, lista, reloj):
    otro = lista.query(PrecioProducto).filter_by(producto="71 CCC3").one()
    antes = otro.actualizado_en

    _cargar(client, _corto(["71 AAA1", "990", "", "", ""]))

    lista.expire_all()
    assert lista.query(PrecioProducto).filter_by(producto="71 CCC3").one().actualizado_en == antes


def test_el_tope_de_cambios_de_la_vista_previa_del_corto_es_de_mil(client, lista, reloj, monkeypatch):
    assert carga.TOPE_CAMBIOS_CORTO == 1000
    monkeypatch.setattr(carga, "TOPE_CAMBIOS_CORTO", 2)

    v = _vista(client, _corto(["71 AAA1", "990", "", "", ""], ["71 CCC3", "2500", "", "", ""],
                              ["13 BBB2", "1500", "", "", ""]))

    assert v["cambios_total"] == 3 and len(v["cambios"]) == 2


def test_el_corto_no_acepta_mas_de_2000_productos(client, lista, reloj):
    sobra = _corto(*[[f"99 X{i}", "100", "", "", ""] for i in range(2001)])
    justo = _corto(*[[f"99 X{i}", "100", "", "", ""] for i in range(2000)])

    r = _post(client, sobra)
    v = _vista(client, justo)

    assert r.status_code == 422 and "2.000" in r.json()["detail"] and "archivo completo" in r.json()["detail"]
    assert v["filas"] == 2000 and v["desconocidos_total"] == 2000 and v["errores_total"] == 0


# ================================================================== 10. permisos
def test_solo_el_equipo_de_precios_puede_cargar_o_bajar_la_plantilla(client, lista):
    contenido = _corto(["71 AAA1", "990", "", "", ""])
    app.dependency_overrides[requiere_auth] = lambda: "noadmin@curifor.com"
    try:
        with TestClient(app) as c:
            assert _post(c, contenido).status_code == 403
            assert _cargar(c, contenido).status_code == 403
            assert c.get("/api/precios/carga/plantilla").status_code == 403
    finally:
        app.dependency_overrides[requiere_auth] = lambda: "test@curifor.com"
    assert _override(lista, "71 AAA1") is None


# ================================================================== 11. xlsx
def _xlsx(filas, cab=CAB) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(list(cab))
    for f in filas:
        ws.append(list(f))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_un_xlsx_con_fechas_y_numeros_de_verdad_se_lee(client, lista, reloj):
    contenido = _xlsx([["71 AAA1", 990, date(2026, 10, 20), date(2026, 10, 31), "campaña"],
                       ["71 CCC3", 2500.0, None, None, None]])

    v = _vista(client, contenido, nombre="corto.xlsx")
    r = _cargar(client, contenido, nombre="corto.xlsx")

    assert v["modo"] == "corto" and v["a_cambiar"] == 2 and v["errores_total"] == 0
    assert r.status_code == 200 and r.json()["aplicados"] == 2
    o = _override(lista, "71 AAA1")
    assert (o.precio_fijo, o.vigente_desde, o.vigente_hasta) == (990, date(2026, 10, 20), date(2026, 10, 31))
    assert _override(lista, "71 CCC3").precio_fijo == 2500


def test_un_xlsx_con_el_codigo_como_numero_se_lee_como_texto(client, lista, reloj):
    _producto(lista, "12345", procedencia_maestro="NACIONAL")
    contenido = _xlsx([[12345, 990, None, None, None]])

    v = _vista(client, contenido, nombre="corto.xlsx")

    assert v["a_cambiar"] == 1 and v["desconocidos_total"] == 0


def test_un_xlsx_cuyos_datos_no_estan_en_la_primera_hoja_lo_dice(client, lista):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Notas"
    wb.active["A1"] = "apuntes"
    datos = wb.create_sheet("Datos")
    datos.append(CAB)
    datos.append(["71 AAA1", 990, None, None, None])
    buf = io.BytesIO()
    wb.save(buf)

    r = _post(client, buf.getvalue(), nombre="corto.xlsx")

    assert r.status_code == 422 and "solo se lee la primera hoja" in r.json()["detail"]


# ================================================================ 12. la plantilla
def _plantilla(client, archivo=None):
    return client.get("/api/precios/carga/plantilla", params={} if archivo is None else {"archivo": archivo})


@pytest.mark.parametrize("archivo,tipo", [("xlsx", "spreadsheetml"), ("csv", "text/csv")])
def test_la_plantilla_se_descarga_y_no_choca_con_la_ruta_de_un_producto(client, lista, archivo, tipo):
    r = _plantilla(client, archivo)

    assert r.status_code == 200, r.text                      # no un 404 "El producto no esta en la lista"
    assert tipo in r.headers["content-type"]
    assert f".{archivo}" in r.headers["content-disposition"] and "plantilla" in r.headers["content-disposition"]


def test_la_plantilla_por_defecto_es_un_xlsx(client, lista):
    r = _plantilla(client)

    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert r.content[:2] == b"PK"


def test_un_formato_de_plantilla_que_no_existe_es_un_422(client, lista):
    for raro in ("pdf", "xls", ""):
        assert _plantilla(client, raro).status_code == 422, raro


def test_la_plantilla_csv_tiene_los_encabezados_con_punto_y_coma_y_bom(client, lista):
    r = _plantilla(client, "csv")

    assert r.content == "﻿Código;Precio fijo;Desde;Hasta;Obs\r\n".encode("utf-8")


def test_la_plantilla_xlsx_trae_solo_encabezados_validaciones_e_instrucciones(client, lista):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(_plantilla(client, "xlsx").content))

    primera = wb.worksheets[0]
    assert [c.value for c in primera[1]] == CAB
    assert primera.max_row == 1                              # sin filas de ejemplo: se cargarian
    assert wb.sheetnames[1] == "Instrucciones" and wb.worksheets[1].max_row > 5
    tipos = {v.type: v for v in primera.data_validations.dataValidation}
    assert set(tipos) == {"date", "whole"}
    assert "D2" in str(tipos["date"].sqref) and "C2" in str(tipos["date"].sqref)
    assert "B2" in str(tipos["whole"].sqref)
    # QUITAR tiene que poder escribirse: la validacion avisa, no prohibe.
    assert all(v.errorStyle == "warning" for v in tipos.values())


def test_las_instrucciones_de_la_plantilla_explican_quitar_y_las_fechas(client, lista):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(_plantilla(client, "xlsx").content))
    texto = "\n".join(str(c.value) for fila in wb["Instrucciones"].iter_rows() for c in fila if c.value)

    for clave in ("QUITAR", "dd-mm-aaaa", "2.000", "vacía", "Precio fijo", "Desde", "Hasta"):
        assert clave in texto, clave


@pytest.mark.parametrize("archivo,nombre", [("xlsx", "plantilla.xlsx"), ("csv", "plantilla.csv")])
def test_la_plantilla_vacia_se_sube_sin_error_y_con_cero_filas(client, lista, archivo, nombre):
    contenido = _plantilla(client, archivo).content

    for aplicar in (False, True):
        r = _post(client, contenido, nombre=nombre, aplicar=aplicar)
        assert r.status_code == 200, r.text
        j = r.json()
        if aplicar:
            assert j["aplicados"] == 0
        else:
            assert (j["modo"], j["filas"], j["a_cambiar"], j["errores_total"]) == ("corto", 0, 0, 0)
            assert "ningún producto" in j["mensaje"]


def test_la_plantilla_llena_se_carga(client, lista, reloj):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(_plantilla(client, "xlsx").content))
    ws = wb.worksheets[0]
    ws.append(["71 AAA1", 990, date(2026, 10, 20), date(2026, 10, 31), "campaña"])
    buf = io.BytesIO()
    wb.save(buf)

    r = _cargar(client, buf.getvalue(), nombre="plantilla.xlsx")

    assert r.status_code == 200, r.text
    assert r.json()["aplicados"] == 1 and _override(lista, "71 AAA1").vigente_hasta == date(2026, 10, 31)


def test_el_aviso_de_archivo_vacio_no_sale_si_hay_filas(client, lista, reloj):
    con = _vista(client, _corto(["71 AAA1", "990", "", "", ""]))
    sin = _vista(client, _corto(["71 AAA1", "", "", "", ""]))

    assert con.get("mensaje") is None and sin.get("mensaje") is None
    assert sin["a_cambiar"] == 0 and sin["sin_cambios"] == 1


# ============================================ 13. entre comparar y escribir (la ventana)
def test_un_producto_que_se_borra_entre_comparar_y_escribir_no_da_500_en_el_corto(client, lista, reloj, monkeypatch):
    real = carga.planificar

    def con_carrera(db, nombre, contenido, modo="auto"):
        plan = real(db, nombre, contenido, modo)
        db.query(PrecioProducto).filter_by(producto="71 AAA1").delete()
        db.commit()
        return plan

    monkeypatch.setattr(carga, "planificar", con_carrera)

    r = _cargar(client, _corto(["71 AAA1", "990", "", "", ""]))

    assert r.status_code == 200
    assert r.json()["aplicados"] == 0 and r.json()["omitidos"]["desconocidos"] == 1


def test_lo_que_otra_persona_deja_entre_comparar_y_escribir_se_lee_otra_vez_en_el_corto(
        client, lista, reloj, monkeypatch):
    """El corto se aplica contra lo que hay AL ESCRIBIR, con las filas bloqueadas: un margen que
    alguien puso despues de la vista previa tambien lo reemplaza el precio fijo."""
    real = carga.planificar

    def con_carrera(db, nombre, contenido, modo="auto"):
        plan = real(db, nombre, contenido, modo)
        db.add(PrecioOverride(tenant_id="curifor", producto="71 AAA1", margen_fijo=35.0))
        db.commit()
        return plan

    monkeypatch.setattr(carga, "planificar", con_carrera)

    r = _cargar(client, _corto(["71 AAA1", "990", "", "", ""]))

    assert r.status_code == 200 and r.json()["aplicados"] == 1
    o = _override(lista, "71 AAA1")
    assert (o.precio_fijo, o.margen_fijo) == (990, None)


def test_un_periodo_que_se_vuelve_invalido_entre_comparar_y_escribir_omite_solo_esa_fila(
        client, lista, reloj, monkeypatch):
    real = carga.planificar

    def con_carrera(db, nombre, contenido, modo="auto"):
        plan = real(db, nombre, contenido, modo)
        # Despues de la vista previa alguien le deja a 71 CCC3 un periodo ya terminado.
        o = db.query(PrecioOverride).filter_by(producto="71 CCC3").one()
        o.vigente_hasta = _dias(-3)
        db.commit()
        return plan

    monkeypatch.setattr(carga, "planificar", con_carrera)

    r = _cargar(client, _corto(["71 CCC3", "2500", "", "", ""], ["71 AAA1", "990", "", "", ""]))

    assert r.status_code == 200
    assert r.json()["aplicados"] == 1 and r.json()["omitidos"]["conflictos"] == 1
    assert _override(lista, "71 AAA1").precio_fijo == 990 and _override(lista, "71 CCC3").precio_fijo is None
